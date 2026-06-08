# -*- coding: utf-8 -*-
"""
Excel to Markdown Converter
============================
Đọc file Excel, mỗi sheet tạo thành 1 file .md riêng biệt.

Output structure:
  output_dir/
    ten-file-excel/
      ten-sheet.md

Usage:
  python excel_to_md.py input.xlsx
  python excel_to_md.py input.xlsx -o ./output
  python excel_to_md.py ./data/*.xlsx -o ./docs/modules

Requirements:
  pip install openpyxl
"""

import argparse
import glob
import io
import os
import re
import sys
import unicodedata
from pathlib import Path

# Force UTF-8 mode on Windows (handles Japanese, Vietnamese filenames)
os.environ['PYTHONUTF8'] = '1'
if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

try:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
except ImportError:
    print('Missing dependency. Install with: pip install openpyxl')
    sys.exit(1)


def sanitize_filename(text: str) -> str:
    """Convert text to a safe filename, preserving Unicode (Japanese, Vietnamese, etc.).

    - Keeps letters, numbers, CJK, Hiragana, Katakana, Vietnamese chars
    - Replaces spaces/underscores with hyphens
    - Removes filesystem-unsafe characters
    """
    text = text.strip()
    # Remove characters unsafe for filesystems: / \ : * ? " < > |
    text = re.sub(r'[\\/:*?"<>|]', '', text)
    # Replace whitespace and underscores with hyphens
    text = re.sub(r'[\s_]+', '-', text)
    # Collapse multiple hyphens
    text = re.sub(r'-+', '-', text)
    return text.strip('-') or 'unnamed'


def get_cell_value(cell) -> str:
    """Get string value from a cell, handling merged cells."""
    if isinstance(cell, MergedCell):
        return ''
    if cell.value is None:
        return ''
    return str(cell.value).strip()


def detect_header_row(ws) -> int | None:
    """Detect which row is the header (first row with multiple non-empty cells)."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(10, ws.max_row or 1)), start=1):
        non_empty = sum(1 for cell in row if get_cell_value(cell))
        if non_empty >= 2:
            return row_idx
    return None


def get_merged_cell_value(ws, row, col) -> str:
    """Resolve value for a cell that might be part of a merged range."""
    for merged_range in ws.merged_cells.ranges:
        if (merged_range.min_row <= row <= merged_range.max_row
                and merged_range.min_col <= col <= merged_range.max_col):
            return get_cell_value(ws.cell(merged_range.min_row, merged_range.min_col))
    return get_cell_value(ws.cell(row, col))


def compute_col_widths(headers: list[str], data_rows: list[list[str]]) -> list[int]:
    """Compute column widths based on content."""
    widths = [len(h) for h in headers]
    for row in data_rows:
        for i, cell in enumerate(row):
            if i < len(widths):
                widths[i] = max(widths[i], len(cell))
    return [max(w, 3) for w in widths]


def sheet_to_markdown(ws) -> str:
    """Convert a worksheet to markdown string."""
    lines: list[str] = []
    sheet_title = ws.title or 'Untitled'
    lines.append(f'# {sheet_title}')
    lines.append('')

    if not ws.max_row or not ws.max_column:
        lines.append('*Sheet is empty.*')
        return '\n'.join(lines)

    header_row_idx = detect_header_row(ws)

    # Collect notes/content above the header row
    if header_row_idx and header_row_idx > 1:
        for row_idx in range(1, header_row_idx):
            row_values = [
                get_merged_cell_value(ws, row_idx, col)
                for col in range(1, ws.max_column + 1)
            ]
            text = ' | '.join(v for v in row_values if v)
            if text:
                lines.append(f'> {text}')
        if lines[-1] != '':
            lines.append('')

    if not header_row_idx:
        # No clear header — dump all rows as a simple table
        header_row_idx = 1

    # Extract headers
    headers = [
        get_merged_cell_value(ws, header_row_idx, col)
        for col in range(1, ws.max_column + 1)
    ]

    # Trim trailing empty columns
    while headers and not headers[-1]:
        headers.pop()

    if not headers:
        lines.append('*Sheet has no data.*')
        return '\n'.join(lines)

    col_count = len(headers)
    headers = [h if h else f'Col {i + 1}' for i, h in enumerate(headers)]

    # Extract data rows
    data_rows: list[list[str]] = []
    for row_idx in range(header_row_idx + 1, (ws.max_row or header_row_idx) + 1):
        row = [
            get_merged_cell_value(ws, row_idx, col)
            for col in range(1, col_count + 1)
        ]
        # Skip completely empty rows
        if any(cell for cell in row):
            data_rows.append(row)

    # Build markdown table
    widths = compute_col_widths(headers, data_rows)

    header_line = '| ' + ' | '.join(h.ljust(w) for h, w in zip(headers, widths)) + ' |'
    sep_line = '| ' + ' | '.join('-' * w for w in widths) + ' |'
    lines.append(header_line)
    lines.append(sep_line)

    for row in data_rows:
        # Pad row to match column count
        padded = row + [''] * (col_count - len(row))
        row_line = '| ' + ' | '.join(
            cell.ljust(w) for cell, w in zip(padded, widths)
        ) + ' |'
        lines.append(row_line)

    lines.append('')

    # Summary
    lines.append(f'---')
    lines.append(f'*Generated from sheet "{sheet_title}" — {len(data_rows)} rows, {col_count} columns.*')

    return '\n'.join(lines)


def resolve_path(file_path: str) -> Path | None:
    """Resolve file path, handling Windows Unicode normalization edge cases."""
    p = Path(file_path)
    if p.exists():
        return p
    # On Windows, Path.exists() can fail when Unicode normalization differs
    # between the input string (e.g. NFD from shell) and filesystem (NFC).
    parent = p.parent
    if not parent.exists():
        return None
    target_nfc = unicodedata.normalize('NFC', p.name)
    for child in parent.iterdir():
        if unicodedata.normalize('NFC', child.name) == target_nfc:
            return child
    return None


def convert_excel(excel_path: str, output_dir: str) -> list[str]:
    """Convert an Excel file to markdown files. Returns list of created file paths."""
    resolved = resolve_path(excel_path)
    if not resolved:
        print(f'  [ERROR] File not found: {excel_path}')
        return []
    excel_path = resolved

    wb = load_workbook(str(excel_path), read_only=False, data_only=True)

    # Folder name = slugified Excel filename (without extension)
    folder_name = sanitize_filename(excel_path.stem)
    if not folder_name:
        folder_name = 'unnamed'

    target_dir = Path(output_dir) / folder_name
    target_dir.mkdir(parents=True, exist_ok=True)

    created_files: list[str] = []

    for ws in wb.worksheets:
        sheet_slug = sanitize_filename(ws.title) if ws.title else 'unnamed-sheet'
        if not sheet_slug:
            sheet_slug = 'unnamed-sheet'

        md_content = sheet_to_markdown(ws)
        md_path = target_dir / f'{sheet_slug}.md'

        # Handle duplicate sheet names
        counter = 1
        while md_path.exists():
            counter += 1
            md_path = target_dir / f'{sheet_slug}-{counter}.md'

        md_path.write_text(md_content, encoding='utf-8')
        created_files.append(str(md_path))
        print(f'  [OK] {ws.title} -> {md_path}')

    wb.close()
    return created_files


def main():
    parser = argparse.ArgumentParser(
        description='Convert Excel files to Markdown. Each sheet becomes a separate .md file.',
        epilog='Example: python excel_to_md.py data.xlsx -o ./output',
    )
    parser.add_argument(
        'inputs',
        nargs='+',
        help='Excel file(s) or glob patterns (e.g., ./data/*.xlsx)',
    )
    parser.add_argument(
        '-o', '--output',
        default='./output',
        help='Output directory (default: ./output)',
    )
    args = parser.parse_args()

    # Expand glob patterns
    excel_files: list[str] = []
    for pattern in args.inputs:
        matches = glob.glob(pattern, recursive=True)
        if matches:
            excel_files.extend(matches)
        else:
            excel_files.append(pattern)

    if not excel_files:
        print('No input files provided.')
        sys.exit(1)

    total_created = 0

    for excel_file in excel_files:
        print(f'\nProcessing: {excel_file}')
        created = convert_excel(excel_file, args.output)
        total_created += len(created)

    print(f'\nDone! Created {total_created} markdown file(s) in: {args.output}/')


if __name__ == '__main__':
    main()
