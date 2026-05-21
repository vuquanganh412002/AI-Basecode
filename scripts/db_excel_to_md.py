#!/usr/bin/env python3
"""
db_excel_to_md.py — Convert VTI-format Database Design Excel → Markdown

Usage:
    python3 db_excel_to_md.py <input.xlsx>
    python3 db_excel_to_md.py <input.xlsx> <output.md>

Output defaults to same directory as input, named database-design.md.
"""

import sys
import re
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is required.\n  pip install openpyxl")
    sys.exit(1)

# ──────────────────────────────────────────────────────────────────────────────
# Column layout (database sheets)
# ──────────────────────────────────────────────────────────────────────────────

# Meta header row 2
DB_HDR = [
    (1,  "system_name"),
    (10, "document_name"),
    (17, "_sheet_name"),
    (24, "created_date"),
    (29, "created_by"),
    (34, "updated_date"),
    (39, "updated_by"),
]

# Table info row 5
COL_PHYS = 10   # physical table name
COL_LOGI = 34   # logical table name

# Column table (row 7 = header, row 8+ = data)
COL_NO      = 1
COL_NAME    = 3
COL_PK      = 8
COL_IX      = [10, 12, 14, 16, 18]   # IX columns (up to 5 indexes)
COL_TYPE    = 20
COL_SIZE    = 24
COL_IDENT   = 27
COL_NULL    = 30
COL_NOTE    = 34

# Index section
IDX_NO    = 1
IDX_NAME  = 3
IDX_COL   = 12
IDX_PK    = 20
IDX_UQ    = 27
IDX_NOTE  = 34

# Cover page
COV_ROW = {
    "format_code":    20,
    "format_version": 21,
    "issue_date":     22,
}
COV_VAL = 23   # value column (W)

NON_TABLE_SHEETS = {"表紙", "変更履歴", "目次", "概要"}


# ──────────────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────────────

def _v(ws, row, col):
    """Return effective merged-cell value as string."""
    cell = ws.cell(row=row, column=col)
    val  = cell.value
    if val is None:
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                val = ws.cell(row=rng.min_row, column=rng.min_col).value
                break
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y/%m/%d")
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).rstrip()


# ──────────────────────────────────────────────────────────────────────────────
# Readers
# ──────────────────────────────────────────────────────────────────────────────

def _v_raw(ws, row, col):
    """Return raw cell value preserving float decimals (used for version fields)."""
    cell = ws.cell(row=row, column=col)
    val  = cell.value
    if val is None:
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                val = ws.cell(row=rng.min_row, column=rng.min_col).value
                break
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y/%m/%d")
    if isinstance(val, float):
        # Preserve decimals (e.g. 1.0 → "1.0", 1.1 → "1.1")
        return f"{val:.1f}" if val == round(val, 1) else str(val)
    return str(val).strip()


def read_cover(wb, input_path):
    data = {}
    if "表紙" in wb.sheetnames:
        ws = wb["表紙"]
        for key, row in COV_ROW.items():
            if key == "format_version":
                data[key] = _v_raw(ws, row, COV_VAL)
            else:
                data[key] = _v(ws, row, COV_VAL)
    m = re.match(r"【(.+?)】", input_path.stem)
    data["customer_name"] = m.group(1) if m else "日本農業新聞様"
    return data


def read_meta(ws):
    return {key: _v(ws, 2, col) for col, key in DB_HDR}


def read_changelog(wb):
    if "変更履歴" not in wb.sheetnames:
        return []
    ws   = wb["変更履歴"]
    rows = []
    r    = 3   # row 1: title, row 2: header
    while r <= ws.max_row:
        no = _v(ws, r, 1)
        if not no or not no.replace(".", "").isdigit():
            break
        date = _v(ws, r, 2)
        if not date:   # skip empty placeholder rows
            r += 1; continue
        rows.append({
            "no":          str(int(float(no))),
            "date":        date,
            "version":     _v_raw(ws, r, 7),
            "author":      _v(ws, r, 11),
            "description": _v(ws, r, 18),
            "reviewer":    _v(ws, r, 31),
            "approver":    _v(ws, r, 38),
        })
        r += 1
    return rows


def read_overview(wb):
    if "概要" not in wb.sheetnames:
        return "", "", []
    ws       = wb["概要"]
    sys_lines = []
    doc_lines = []
    related   = []
    state     = None

    r = 5
    while r <= ws.max_row:
        b = _v(ws, r, 2)
        c = _v(ws, r, 3)

        if "システム概要" in b:
            state = "sys"; r += 1; continue
        elif "資料概要" in b:
            state = "doc"; r += 1; continue
        elif "関連資料" in b:
            state = "related"; r += 1; continue

        if state == "sys":
            sys_lines.append(c)   # keep empty lines to preserve paragraph breaks
        elif state == "doc":
            doc_lines.append(c)
        elif state == "related":
            # skip header row
            no_val = _v(ws, r, 3)
            if no_val == "No":
                r += 1; continue
            if no_val and no_val.replace(".", "").isdigit():
                code = _v(ws, r, 4)
                if not code:   # skip empty placeholder rows
                    r += 1; continue
                related.append({
                    "no":   str(int(float(no_val))),
                    "code": code,
                    "name": _v(ws, r, 15),
                })

        r += 1

    # Strip trailing empty lines, join with newlines
    def _trim_join(lines):
        while lines and not lines[-1]:
            lines.pop()
        return "\n".join(lines)

    return _trim_join(sys_lines), _trim_join(doc_lines), related


def read_table_sheet(ws):
    """Read a database table sheet → dict."""
    meta = read_meta(ws)

    physical = _v(ws, 5, COL_PHYS)
    logical  = _v(ws, 5, COL_LOGI)

    # Find header row (contains "No" in col 1 and "項目名" in col 3)
    header_row = None
    for r in range(6, min(15, ws.max_row + 1)):
        if _v(ws, r, COL_NO) == "No" and _v(ws, r, COL_NAME) == "項目名":
            header_row = r
            break
    if header_row is None:
        return None

    # Read columns
    columns = []
    r = header_row + 1
    while r <= ws.max_row:
        no_val = _v(ws, r, COL_NO)
        if not no_val or not no_val.replace(".", "").isdigit():
            break
        no_int = str(int(float(no_val)))
        columns.append({
            "no":       no_int,
            "name":     _v(ws, r, COL_NAME),
            "pk":       _v(ws, r, COL_PK),
            "type":     _v(ws, r, COL_TYPE),
            "size":     _v(ws, r, COL_SIZE),
            "identity": _v(ws, r, COL_IDENT),
            "nullable": _v(ws, r, COL_NULL),
            "note":     _v(ws, r, COL_NOTE),
        })
        r += 1

    # Read indexes (look for インデックス情報 label)
    indexes = []
    idx_header = None
    while r <= ws.max_row:
        a = _v(ws, r, 1)
        if "インデックス" in a:
            idx_header = r + 1
            break
        r += 1

    if idx_header:
        r = idx_header + 1   # skip header row
        while r <= ws.max_row:
            no_val = _v(ws, r, IDX_NO)
            if not no_val or not no_val.replace(".", "").isdigit():
                break
            indexes.append({
                "no":       str(int(float(no_val))),
                "name":     _v(ws, r, IDX_NAME),
                "columns":  _v(ws, r, IDX_COL).strip().rstrip("\t"),
                "pk":       _v(ws, r, IDX_PK),
                "unique":   _v(ws, r, IDX_UQ),
                "note":     _v(ws, r, IDX_NOTE),
            })
            r += 1

    return {
        "physical":  physical,
        "logical":   logical,
        "meta":      meta,
        "columns":   columns,
        "indexes":   indexes,
    }


# ──────────────────────────────────────────────────────────────────────────────
# Markdown builder
# ──────────────────────────────────────────────────────────────────────────────

def build_md(wb, input_path):
    cover = read_cover(wb, input_path)

    # Use meta from first table sheet for dates/author
    meta = {}
    for s in wb.sheetnames:
        if s not in NON_TABLE_SHEETS and s != "目次":
            meta = read_meta(wb[s])
            break

    changelog = read_changelog(wb)
    sys_text, doc_text, related = read_overview(wb)

    tables = []
    for s in wb.sheetnames:
        if s in NON_TABLE_SHEETS or s == "目次":
            continue
        t = read_table_sheet(wb[s])
        if t:
            tables.append(t)

    L = []

    # ── Frontmatter ────────────────────────────────────────────────────────────
    fmt_ver = cover.get("format_version", "1.0")
    # Normalize integer versions: "1" → "1.0", "2" → "2.0"
    if fmt_ver and re.match(r"^\d+$", fmt_ver):
        fmt_ver = fmt_ver + ".0"
    issue   = cover.get("issue_date", "").replace("/", "-")
    # issue_date may come as "2/22/2019" → normalize
    try:
        from datetime import datetime as _dt
        for fmt in ("%m/%d/%Y", "%Y/%m/%d", "%Y-%m-%d"):
            try:
                issue = _dt.strptime(issue.replace("-", "/"), fmt.replace("-", "/")).strftime("%Y-%m-%d")
                break
            except ValueError:
                pass
    except Exception:
        pass

    L += [
        "---",
        f"customer_name: {cover.get('customer_name', '日本農業新聞様')}",
        f"system_name: {meta.get('system_name', cover.get('system_name', ''))}",
        f"document_name: {meta.get('document_name', cover.get('document_name', ''))}",
        f"format_code: {cover.get('format_code', '17-BM/PM/VTI')}",
        f'format_version: "{fmt_ver}"',
        f"issue_date: {issue}",
        f"created_date: {meta.get('created_date', '')}",
        f"created_by: {meta.get('created_by', '')}",
        f"updated_date: {meta.get('updated_date', '')}",
        f"updated_by: {meta.get('updated_by', '')}",
        "---", "",
    ]

    # ── 変更履歴 ──────────────────────────────────────────────────────────────
    L += [
        "## 変更履歴", "",
        "| No | 発行日 | バージョン | 担当者 | 変更内容 | 確認者 | 承認者 |",
        "|---|---|---|---|---|---|---|",
    ]
    for ch in changelog:
        L.append(f"| {ch['no']} | {ch['date']} | {ch['version']} | {ch['author']} | {ch['description']} | {ch['reviewer']} | {ch['approver']} |")
    L.append("")

    # ── システム概要 ──────────────────────────────────────────────────────────
    L += ["## システム概要", "", sys_text, ""]

    # ── 資料概要 ──────────────────────────────────────────────────────────────
    L += ["## 資料概要", "", doc_text, ""]

    # ── 関連資料 ──────────────────────────────────────────────────────────────
    L += [
        "## 関連資料", "",
        "| No | 資料コード | 資料名 |",
        "|---|---|---|",
    ]
    for rd in related:
        L.append(f"| {rd['no']} | {rd['code']} | {rd['name']} |")
    L.append("")

    def _cell(v):
        """Escape newlines in table cell values so MD rows stay single-line."""
        return str(v).replace("\n", "<br>")

    # ── Table sections ────────────────────────────────────────────────────────
    for t in tables:
        L += ["---", "", f"# {t['physical']} ({t['logical']})", ""]

        L += [
            "| No | 項目名 | PK | 属性 | サイズ | IDENTITY | NULL許容 | 備考 |",
            "|---|---|---|---|---|---|---|---|",
        ]
        for c in t["columns"]:
            L.append(
                f"| {_cell(c['no'])} | {_cell(c['name'])} | {_cell(c['pk'])} | {_cell(c['type'])} "
                f"| {_cell(c['size'])} | {_cell(c['identity'])} | {_cell(c['nullable'])} | {_cell(c['note'])} |"
            )
        L.append("")

        if t["indexes"]:
            L += [
                "## インデックス", "",
                "| 項番 | インデックス名 | カラム名 | 主キー | ユニーク | 備考 |",
                "|---|---|---|---|---|---|",
            ]
            for ix in t["indexes"]:
                L.append(
                    f"| {_cell(ix['no'])} | {_cell(ix['name'])} | {_cell(ix['columns'])} "
                    f"| {_cell(ix['pk'])} | {_cell(ix['unique'])} | {_cell(ix['note'])} |"
                )
            L.append("")

    return "\n".join(L)


# ──────────────────────────────────────────────────────────────────────────────
# Entry point
# ──────────────────────────────────────────────────────────────────────────────

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Convert VTI Database Excel to Markdown")
    parser.add_argument("input",  help="Input .xlsx file")
    parser.add_argument("output", nargs="?", help="Output .md file (optional)")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}")
        sys.exit(1)

    print(f"Reading: {input_path.name} ...")
    wb = openpyxl.load_workbook(str(input_path), data_only=True)

    md = build_md(wb, input_path)

    if args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.parent / "database-design.md"

    output_path.write_text(md, encoding="utf-8")
    print(f"✓  Saved: {output_path}")


if __name__ == "__main__":
    main()
