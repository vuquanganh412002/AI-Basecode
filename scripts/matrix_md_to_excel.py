#!/usr/bin/env python3
"""
matrix_md_to_excel.py — Convert a design-note Markdown file (tables + prose)
to a styled .xlsx workbook.

Written for docs/design-vi/dokusya-change-matrix.md — a 27-column change
matrix that is unreadable in a terminal — but the parser is generic: any
Markdown made of headings, GFM tables, paragraphs, blockquotes, lists and
fenced code blocks converts cleanly.

Usage:
    python3 matrix_md_to_excel.py <input.md>
    python3 matrix_md_to_excel.py <input.md> <output.xlsx>
    python3 matrix_md_to_excel.py <input.md> --split-level 3
    python3 matrix_md_to_excel.py <input.md> --no-color

Default output: same directory as <input.md>, same basename, `.xlsx`.

Sheet layout:
    - One sheet per heading at `--split-level` (default 2 = every `##`).
      Content above the first split heading goes to a leading "Tổng quan"
      sheet together with the source path + generation timestamp.
    - Wide tables (>= 10 columns) get frozen panes on the header row and
      the first column, so the field-name column stays visible while
      scrolling right through the mode columns.

Cell rendering:
    - `**bold**` and `` `code` `` survive as rich text (Excel rich-text runs),
      because in the change matrix bold marks "this value actually changed".
    - Legend glyphs are colour-coded when --no-color is not passed:
        ✓ ✅  green      ✗ ❌  red      ⧗  amber
        — ↩   muted grey (unchanged / inherited)
    - Rows whose first cell looks like `━━ section ━━` become full-width
      section bands.

Sibling of scripts/api_md_to_excel.py and scripts/testcase_md_to_excel.py —
style constants are kept in sync with those for visual coherence.
"""
from __future__ import annotations

import argparse
import math
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("ERROR: openpyxl is required.\n  pip install openpyxl")
    sys.exit(1)

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    RICH_TEXT = True
except ImportError:  # openpyxl < 3.1 — degrade to plain strings
    RICH_TEXT = False


# ─── Style constants (match scripts/api_md_to_excel.py) ──────────────────

HEADER_BLUE = "FFB4C6E7"   # table column-header fill
SUB_HEADER = "FFBDD6EE"    # section-band rows inside a table
LIGHT_BLUE = "FFDEEAF6"    # heading rows
COVER_META = "FFD9E2F3"    # overview metadata box
NAVY = "FF1F3864"          # heading text colour
CODE_BG = "FFF2F2F2"       # fenced code block background

OK_BG = "FFE2EFDA"         # ✓ / ✅
OK_FG = "FF1E6C41"
NG_BG = "FFFCE4E4"         # ✗ / ❌
NG_FG = "FFB00020"
WAIT_BG = "FFFFF2CC"       # ⧗ (deferred until the batch runs)
WAIT_FG = "FF7F6000"
MUTED_FG = "FF9A9A9A"      # — / ↩

FONT_NAME = "游ゴシック"
CODE_FONT = "Consolas"

BODY_SIZE = 9
HEAD_SIZE = 10
PROSE_SIZE = 10

WIDE_TABLE_COLS = 10       # >= this many columns ⇒ freeze panes + centre body
PROSE_SPAN = 8             # columns a paragraph is merged across
MIN_WIDTH, MAX_WIDTH = 8.0, 44.0

_THIN = Side(style="thin", color="FFBFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold: bool = False, size: int = BODY_SIZE, color: str = "FF000000",
          name: str = FONT_NAME, italic: bool = False) -> Font:
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)


def _align(wrap: bool = True, h: str = "left", v: str = "center") -> Alignment:
    return Alignment(wrap_text=wrap, horizontal=h, vertical=v)


# ─── Markdown parsing ────────────────────────────────────────────────────

@dataclass
class Block:
    """One Markdown construct. `kind` drives how it is rendered."""
    kind: str                                   # heading|table|para|quote|list|code|hr
    text: str = ""
    level: int = 0                              # heading level
    rows: list[list[str]] = field(default_factory=list)   # table
    lines: list[str] = field(default_factory=list)        # code / list


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_FENCE_RE = re.compile(r"^\s*```+\s*(\S*)\s*$")
_HR_RE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")
_LIST_RE = re.compile(r"^\s*([-*+]|\d+\.)\s+")


def _split_row(line: str) -> list[str]:
    """Split one Markdown table row into cells, honouring `\\|` escapes."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    para: list[str] = []
    i = 0

    def flush_para() -> None:
        nonlocal para
        if para:
            blocks.append(Block("para", text=" ".join(s.strip() for s in para)))
            para = []

    while i < len(lines):
        line = lines[i]

        fence = _FENCE_RE.match(line)
        if fence:
            flush_para()
            lang = fence.group(1)
            i += 1
            code: list[str] = []
            while i < len(lines) and not _FENCE_RE.match(lines[i]):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence
            blocks.append(Block("code", text=lang, lines=code))
            continue

        if not line.strip():
            flush_para()
            i += 1
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            flush_para()
            blocks.append(Block("heading", text=heading.group(2).strip(),
                                level=len(heading.group(1))))
            i += 1
            continue

        if _HR_RE.match(line):
            flush_para()
            blocks.append(Block("hr"))
            i += 1
            continue

        # Table: a `|` row immediately followed by a `|---|` separator.
        if line.lstrip().startswith("|") and i + 1 < len(lines) \
                and _TABLE_SEP_RE.match(lines[i + 1]):
            flush_para()
            rows = [_split_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            for r in rows:
                r.extend([""] * (width - len(r)))
            blocks.append(Block("table", rows=rows))
            continue

        if line.lstrip().startswith(">"):
            flush_para()
            quote: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(lines[i].lstrip()[1:].strip())
                i += 1
            blocks.append(Block("quote", text=" ".join(q for q in quote if q)))
            continue

        if _LIST_RE.match(line):
            flush_para()
            items: list[str] = []
            while i < len(lines) and _LIST_RE.match(lines[i]):
                items.append(_LIST_RE.sub("", lines[i]).strip())
                i += 1
            blocks.append(Block("list", lines=items))
            continue

        para.append(line)
        i += 1

    flush_para()
    return blocks


# ─── Inline formatting → Excel rich text ─────────────────────────────────

_INLINE_RE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"          # **bold**
    r"|`(?P<code>[^`]+)`"             # `code`
    r"|\[(?P<link>[^\]]+)\]\([^)]*\)"  # [text](url) → text
)


def _segments(text: str, bold: bool = False) -> list[tuple[str, bool, bool]]:
    """Split inline Markdown into (text, bold, monospace) runs.

    Recurses into `**bold**` so that a code span nested inside bold
    (`**\\`zougen_hokoku_flg\\`**`, common in this doc) keeps both traits
    instead of leaking literal backticks into the cell.
    """
    out: list[tuple[str, bool, bool]] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], bold, False))
        if m.group("bold") is not None:
            out.extend(_segments(m.group("bold"), bold=True))
        elif m.group("code") is not None:
            out.append((m.group("code"), bold, True))
        else:
            out.append((m.group("link"), bold, False))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], bold, False))
    return [s for s in out if s[0]]


def plain_text(text: str) -> str:
    return "".join(seg[0] for seg in _segments(text))


def _cell_value(text: str, size: int, base_bold: bool):
    """Return a str, or CellRichText when the cell mixes formatting runs."""
    segs = _segments(text)
    if not segs:
        return ""
    if not RICH_TEXT or all(not b and not c for _, b, c in segs):
        return "".join(s for s, _, _ in segs)
    blocks = []
    for chunk, bold, code in segs:
        blocks.append(TextBlock(
            InlineFont(rFont=CODE_FONT if code else FONT_NAME,
                       sz=size, b=bold or base_bold),
            chunk,
        ))
    return CellRichText(*blocks)


# ─── Legend-glyph colouring ──────────────────────────────────────────────

_BAND_RE = re.compile(r"^\s*━+.*━+\s*$")


def _cell_style(value: str) -> tuple[str | None, str | None]:
    """(fill, font colour) for a matrix cell, from its legend glyphs."""
    v = value.strip()
    if not v:
        return None, None
    if "✗" in v or "❌" in v:
        return NG_BG, NG_FG
    if "⧗" in v:
        return WAIT_BG, WAIT_FG
    if "✓" in v or "✅" in v:
        return OK_BG, OK_FG
    if "⚠" in v:
        return WAIT_BG, WAIT_FG
    if v in {"—", "↩", "-"}:
        return None, MUTED_FG
    return None, None


# ─── Workbook rendering ──────────────────────────────────────────────────

def _visual_len(text: str) -> int:
    """Approximate display width — CJK / full-width glyphs count double."""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


class SheetWriter:
    """Renders a list of Blocks onto one worksheet."""

    def __init__(self, ws, colorize: bool) -> None:
        self.ws = ws
        self.colorize = colorize
        self.row = 1
        self.widths: dict[int, float] = {}
        self.frozen = False

    # -- helpers ---------------------------------------------------------

    def _want_width(self, col: int, text: str, cap: float = MAX_WIDTH) -> None:
        want = min(cap, max(MIN_WIDTH, _visual_len(text) * 0.72 + 2))
        self.widths[col] = max(self.widths.get(col, 0.0), want)

    def _merged_line(self, text: str, *, font: Font, fill: str | None = None,
                     height: float | None = None, indent_rows: int = 0) -> None:
        for _ in range(indent_rows):
            self.row += 1
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1,
                       end_row=self.row, end_column=PROSE_SPAN)
        cell = ws.cell(row=self.row, column=1, value=_cell_value(text, font.size, font.bold))
        cell.font = font
        cell.alignment = _align(wrap=True, h="left", v="top")
        if fill:
            for col in range(1, PROSE_SPAN + 1):
                ws.cell(row=self.row, column=col).fill = _fill(fill)
        if height:
            ws.row_dimensions[self.row].height = height
        self.row += 1

    def _prose_height(self, text: str) -> float:
        # ~110 visual units fit across the merged span at PROSE_SIZE.
        return max(16.0, math.ceil(_visual_len(text) / 110) * 14.5)

    # -- block renderers -------------------------------------------------

    def heading(self, block: Block) -> None:
        size = {1: 14, 2: 12, 3: 11}.get(block.level, 10)
        if self.row > 1:
            self.row += 1
        self._merged_line(block.text, font=_font(True, size, NAVY),
                          fill=LIGHT_BLUE if block.level <= 3 else None,
                          height=size + 8)

    def paragraph(self, block: Block, *, italic: bool = False,
                  bullet: str = "") -> None:
        text = f"{bullet}{block.text}" if bullet else block.text
        self._merged_line(text,
                          font=_font(False, PROSE_SIZE, italic=italic),
                          height=self._prose_height(text))

    def quote(self, block: Block) -> None:
        self._merged_line(block.text,
                          font=_font(False, PROSE_SIZE, "FF44546A", italic=True),
                          fill=COVER_META,
                          height=self._prose_height(block.text))

    def bullets(self, block: Block) -> None:
        for item in block.lines:
            text = f"・{item}"
            self._merged_line(text, font=_font(False, PROSE_SIZE),
                              height=self._prose_height(text))

    def code(self, block: Block) -> None:
        for line in block.lines:
            ws = self.ws
            ws.merge_cells(start_row=self.row, start_column=1,
                           end_row=self.row, end_column=PROSE_SPAN)
            cell = ws.cell(row=self.row, column=1, value=line)
            cell.font = _font(False, PROSE_SIZE, name=CODE_FONT)
            cell.alignment = _align(wrap=False, h="left", v="center")
            for col in range(1, PROSE_SPAN + 1):
                ws.cell(row=self.row, column=col).fill = _fill(CODE_BG)
            ws.row_dimensions[self.row].height = 14
            self.row += 1

    def table(self, block: Block) -> None:
        ws = self.ws
        rows = block.rows
        header, body = rows[0], rows[1:]
        ncols = len(header)
        wide = ncols >= WIDE_TABLE_COLS
        header_row = self.row

        for col, text in enumerate(header, start=1):
            cell = ws.cell(row=header_row, column=col,
                           value=_cell_value(text, HEAD_SIZE, True))
            cell.font = _font(True, HEAD_SIZE)
            cell.fill = _fill(HEADER_BLUE)
            cell.alignment = _align(wrap=True, h="center", v="center")
            cell.border = _BORDER
            self._want_width(col, plain_text(text), 46.0 if col == 1 else MAX_WIDTH)
        ws.row_dimensions[header_row].height = 26 if wide else 20
        self.row += 1

        for raw in body:
            first_plain = plain_text(raw[0])
            band = bool(_BAND_RE.match(first_plain))
            for col, text in enumerate(raw, start=1):
                plain = plain_text(text)
                cell = ws.cell(row=self.row, column=col,
                               value=_cell_value(text, BODY_SIZE, band))
                cell.border = _BORDER
                cell.alignment = _align(
                    wrap=True,
                    h="left" if (col == 1 or not wide) else "center",
                    v="center",
                )
                fill_hex, font_hex = (None, None)
                if self.colorize and not band:
                    fill_hex, font_hex = _cell_style(plain)
                if band:
                    cell.fill = _fill(SUB_HEADER)
                elif fill_hex:
                    cell.fill = _fill(fill_hex)
                # Emphasis comes from the source's own `**bold**` runs — the
                # first column is not force-bolded, so rich-text and plain
                # cells render consistently.
                cell.font = _font(band, BODY_SIZE, font_hex or "FF000000")
                self._want_width(col, plain, 46.0 if col == 1 else 22.0)
            self.row += 1

        if wide and not self.frozen:
            ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
            self.frozen = True
        self.row += 1

    # -- finish ----------------------------------------------------------

    def finalize(self) -> None:
        for col, width in self.widths.items():
            self.ws.column_dimensions[get_column_letter(col)].width = width
        for col in range(1, PROSE_SPAN + 1):
            letter = get_column_letter(col)
            if self.ws.column_dimensions[letter].width in (None, 0):
                self.ws.column_dimensions[letter].width = 18.0
        self.ws.sheet_view.zoomScale = 85
        self.ws.sheet_view.showGridLines = False


_INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(title: str, used: set[str]) -> str:
    name = _INVALID_SHEET.sub("-", plain_text(title)).strip() or "Sheet"
    name = name[:31].rstrip(" ,-—·（(")  # don't end a truncated name mid-punctuation
    name = name or "Sheet"
    base, n = name, 2
    while name in used:
        suffix = f" ({n})"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def build_workbook(blocks: list[Block], source: Path, split_level: int,
                   colorize: bool):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set[str] = set()

    doc_title = next((b.text for b in blocks if b.kind == "heading" and b.level == 1),
                     source.stem)

    # Blocks before the first split heading form the overview sheet. Any
    # heading from level 2 down to `split_level` starts a new sheet, so
    # --split-level 3 breaks on `##` AND `###`; level 1 keeps one sheet.
    sections: list[tuple[str, list[Block]]] = [("Tổng quan", [])]
    for block in blocks:
        if block.kind == "heading" and 2 <= block.level <= split_level:
            sections.append((block.text, [block]))
        else:
            sections[-1][1].append(block)

    for title, section in sections:
        if not any(b.kind != "hr" for b in section) and title == "Tổng quan":
            continue
        ws = wb.create_sheet(_sheet_name(title, used))
        writer = SheetWriter(ws, colorize)

        if title == "Tổng quan":
            writer._merged_line(
                plain_text(doc_title),
                font=_font(True, 16, NAVY), fill=COVER_META, height=30)
            stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
            writer._merged_line(
                f"Nguồn: {source.as_posix()}   ·   Xuất lúc: {stamp}",
                font=_font(False, 9, "FF808080"), height=16)
            writer.row += 1

        for block in section:
            if block.kind == "heading":
                if block.level == 1:
                    continue
                writer.heading(block)
            elif block.kind == "table":
                writer.table(block)
            elif block.kind == "quote":
                writer.quote(block)
            elif block.kind == "list":
                writer.bullets(block)
            elif block.kind == "code":
                writer.code(block)
            elif block.kind == "para":
                writer.paragraph(block)
            # 'hr' renders as the blank row a heading already inserts

        writer.finalize()

    return wb


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Convert a design-note Markdown file to a styled .xlsx workbook.")
    parser.add_argument("input", help="source .md file")
    parser.add_argument("output", nargs="?", help="target .xlsx (default: alongside input)")
    parser.add_argument("--split-level", type=int, default=2, choices=(1, 2, 3),
                        help="deepest heading level that starts a new sheet: "
                             "1 = single sheet, 2 = per '##' (default), "
                             "3 = per '##' and '###'")
    parser.add_argument("--no-color", action="store_true",
                        help="skip ✓/✗/⧗ cell colouring")
    args = parser.parse_args(argv)

    src = Path(args.input)
    if not src.is_file():
        print(f"ERROR: not a file: {src}", file=sys.stderr)
        return 2

    dest = Path(args.output) if args.output else src.with_suffix(".xlsx")
    blocks = parse_markdown(src.read_text(encoding="utf-8"))
    if not blocks:
        print(f"ERROR: no content parsed from {src}", file=sys.stderr)
        return 1

    wb = build_workbook(blocks, src, args.split_level, not args.no_color)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)

    tables = sum(1 for b in blocks if b.kind == "table")
    print(f"OK  {dest}")
    print(f"    {len(wb.sheetnames)} sheet(s): {', '.join(wb.sheetnames)}")
    print(f"    {tables} table(s) converted")
    return 0


if __name__ == "__main__":
    sys.exit(main())
