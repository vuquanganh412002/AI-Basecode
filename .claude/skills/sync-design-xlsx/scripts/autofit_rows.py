# -*- coding: utf-8 -*-
"""Recompute row heights for rows whose content changed.

Excel does not auto-fit a row that contains merged cells, and the body cells in
these design documents are merged spans with wrap_text on — so a row whose text
grew keeps its old height and clips. Heights are therefore computed explicitly.

The model is font-aware: Excel column-width units are expressed in the workbook
default font (Calibri 11), so a 6pt marker fits proportionally more characters
per line than 9pt Meiryo body text. It is also East-Asian-width aware, since a
full-width kana costs two units.

Usage:
    from autofit_rows import autofit
    autofit(ws, rows=[13, 14, 17], floor=28.5)
"""
import math
import unicodedata

from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import column_index_from_string

BASE_PT = 11.0          # workbook default font — defines the column width unit
DEFAULT_COL_W = 8.43


def vwidth(s):
    """Visual width in column-width units (full-width chars cost 2)."""
    return sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in s)


def col_widths(ws):
    """Per-column width map, expanding grouped ``<col min=.. max=..>`` entries.

    ``ws.column_dimensions`` registers a grouped definition under its FIRST
    column's letter only. Asking for any other column in that group
    auto-creates a fresh ColumnDimension carrying openpyxl's 13.0 default, so a
    sheet declaring ``<col min="1" max="37" width="5.0"/>`` reports 13.0 for
    columns B..AK — a merged D:AK body cell then measures 442 units instead of
    170 and every height computed from it comes out 2-3x too short.

    Returns ``(widths, default)`` where `widths` is {column_index: width} and
    `default` is the sheet's defaultColWidth for columns no <col> covers.
    """
    default = getattr(ws.sheet_format, 'defaultColWidth', None) or DEFAULT_COL_W
    widths = {}
    for dim in list(ws.column_dimensions.values()):
        if dim.width is None:
            continue
        if dim.min is None:
            # Auto-created on lookup (no XML backing) — not a real declaration.
            if not dim.customWidth:
                continue
            lo = hi = column_index_from_string(dim.index)
        else:
            lo, hi = dim.min, dim.max or dim.min
        for c in range(lo, hi + 1):
            widths[c] = dim.width
    return widths, float(default)


def plain(v):
    """Flatten a cell value to one text stream.

    Rich text MUST be flattened before splitting on newlines: iterating runs and
    rounding each one up to a whole line counts every continuation run on the
    same visual line as a full extra line, which inflates a multi-run cell to
    absurd heights (one row here computed 483pt before this was fixed).
    """
    if isinstance(v, CellRichText):
        return ''.join(b if isinstance(b, str) else b.text for b in v)
    return '' if v is None else str(v)


def autofit(ws, rows, floor=15.0, dry_run=False, ignore_cols=()):
    """Set the height of each row in `rows` to fit its wrapped content.

    `floor` is the section's normal body height — a row never shrinks below it.

    `ignore_cols` are column indices whose content must not drive the height.
    Pass ``(1,)`` on sheets where column A is the VTI marker margin: it is ~5
    units wide, so a row carrying markers from two passes wants four 6pt lines
    and would drag a one-line body row to 39pt. The marker is an annotation in
    the margin — it must never re-lay-out the content.

    Returns {row: new_height} for the rows that actually changed.
    """
    widths, default_w = col_widths(ws)

    merges = {}
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                merges[(r, c)] = (m.min_col, m.max_col)

    changed = {}
    for row in rows:
        # Grow only. The customer hand-sizes rows; a computed height below the
        # current one would re-lay-out parts of the document this sync never
        # touched, so the current height is a floor of its own.
        best = max(floor, ws.row_dimensions[row].height or 0)
        for cell in ws[row]:
            if cell.column in ignore_cols:
                continue
            c1, c2 = merges.get((row, cell.column), (cell.column, cell.column))
            if c1 != cell.column:
                continue                     # merge continuation cell
            if not cell.alignment.wrap_text:
                continue
            text = plain(cell.value)
            if not text:
                continue
            width = sum(widths.get(c, default_w) for c in range(c1, c2 + 1))
            size = cell.font.size or 9.0
            cap = max((width - 1) * (BASE_PT / size), 4)
            lines = sum(max(1, math.ceil(vwidth(seg) / cap))
                        for seg in text.split('\n'))
            best = max(best, lines * (size * 1.5) + 3)
        best = round(best, 2)
        if abs((ws.row_dimensions[row].height or 0) - best) > 0.5:
            changed[row] = best
        if not dry_run:
            ws.row_dimensions[row].height = best
    return changed


if __name__ == '__main__':
    import sys
    from openpyxl import load_workbook
    path, sheet = sys.argv[1], sys.argv[2]
    rows = [int(x) for x in sys.argv[3].split(',')]
    floor = float(sys.argv[4]) if len(sys.argv) > 4 else 15.0
    wb = load_workbook(path, rich_text=True)
    for row, h in sorted(autofit(wb[sheet], rows, floor).items()):
        print(f'{sheet} row {row} -> {h}')
    wb.save(path)
    print('saved — now run restore_cols.py')
