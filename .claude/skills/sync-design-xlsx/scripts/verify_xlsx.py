# -*- coding: utf-8 -*-
"""Post-edit integrity check for a marked-up design workbook.

Compares the edited file against the untouched backup and reports anything the
openpyxl round-trip may have damaged. Run this as the last step of every pass.

Usage: python3 verify_xlsx.py ORIGINAL_BACKUP.xlsx EDITED.xlsx [--mark '8/18 VTI追記']
                              [--color FFE69138]
"""
import os
import re
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_markup import ACCENT, ADDITION, LEGACY_ACCENTS  # noqa: E402

RE_COLS = re.compile(r'<cols>.*?</cols>', re.S)
RE_FMT = re.compile(r'<sheetFormatPr\b[^>]*/>')
# Attributes carrying a namespace prefix (x14ac:dyDescent and friends). Excel
# writes them and declares the prefix on <worksheet>; openpyxl's output root
# declares no such prefix, so restore_cols.py strips them to avoid an unbound-
# prefix parse error. They are cosmetic, so ignore them when comparing.
RE_PREFIXED_ATTR = re.compile(r'\s+[A-Za-z_][\w.-]*:[\w.-]+="[^"]*"')
RE_SHEET = re.compile(r'xl/worksheets/sheet\d+\.xml$')


def _block(xml, rx):
    m = rx.search(xml)
    return m.group(0) if m else None


def main(orig, cur, mark_text, colors):
    zo, zc = zipfile.ZipFile(orig), zipfile.ZipFile(cur)
    fail = []

    print('== layout ==')
    for n in sorted(x for x in zo.namelist() if RE_SHEET.match(x)):
        if n not in zc.namelist():
            fail.append(f'{n} missing'); continue
        a, b = zo.read(n).decode(), zc.read(n).decode()
        same_cols = _block(a, RE_COLS) == _block(b, RE_COLS)
        same_fmt = (RE_PREFIXED_ATTR.sub('', _block(a, RE_FMT))
                    == RE_PREFIXED_ATTR.sub('', _block(b, RE_FMT)))
        if not (same_cols and same_fmt):
            fail.append(f'{n}: cols/sheetFormatPr changed — run restore_cols.py')
        print(f'  {n}: cols={"OK" if same_cols else "DIFF"} '
              f'fmt={"OK" if same_fmt else "DIFF"}')

    print('\n== archive ==')
    bad = zc.testzip()
    print('  zip:', bad or 'OK')
    if bad:
        fail.append(f'corrupt entry {bad}')
    mo = sorted(n for n in zo.namelist() if n.startswith('xl/media/'))
    mc = sorted(n for n in zc.namelist() if n.startswith('xl/media/'))
    print(f'  media: {len(mc)}/{len(mo)} preserved')
    if len(mc) < len(mo):
        fail.append(f'lost media: {set(mo) - set(mc)}')

    print('\n== markup ==')
    wb = load_workbook(cur, rich_text=True)
    strike = rev_run = rev_struck = 0
    marked = {}
    for ws in wb.worksheets:
        rs = [str(r) for r in ws.merged_cells.ranges]
        if len(rs) != len(set(rs)):
            fail.append(f'{ws.title}: duplicate merged ranges')
        rows = set()
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, CellRichText):
                    for b in v:
                        f = getattr(b, 'font', None)
                        if not f:
                            continue
                        is_rev = (f.color is not None
                                  and str(getattr(f.color, 'rgb', '')) in colors)
                        if f.strike:
                            strike += 1
                            # coloured + strike = an earlier revision's
                            # addition, superseded by this one — not a new
                            # correction
                            if is_rev:
                                rev_struck += 1
                        elif is_rev:
                            rev_run += 1
                if c.column == 1 and v is not None and mark_text in str(v):
                    rows.add(c.row)
        if rows:
            marked[ws.title] = sorted(rows)
    print(f'  strikethrough runs: {strike}'
          + (f' ({rev_struck} superseding an earlier revision)' if rev_struck else ''))
    print(f'  revision-coloured runs (rich text): {rev_run}')
    for sheet, rows in marked.items():
        print(f'  marked "{mark_text}" — {sheet}: {len(rows)} rows {rows}')
    print(f'  total marked rows: {sum(len(v) for v in marked.values())}')
    print('  merges: no duplicates' if not any('duplicate' in f for f in fail)
          else '  merges: DUPLICATES FOUND')

    print(f'\nsize: {os.path.getsize(cur):,} bytes '
          f'(original {os.path.getsize(orig):,})')
    if fail:
        print('\n*** FAILED ***')
        for f in fail:
            print('  -', f)
        return 1
    print('\nALL CHECKS PASSED')
    return 0


if __name__ == '__main__':
    flags = {'--mark', '--color'}
    args, skip = [], False
    for a in sys.argv[1:]:
        if skip:
            skip = False
        elif a in flags:
            skip = True
        elif not a.startswith('--'):
            args.append(a)
    mark = '追記'
    # Legacy colours count too: a workbook carrying markup from before the
    # colour change is still a valid, fully-marked-up document.
    colors = {ACCENT, ADDITION, *LEGACY_ACCENTS}
    for i, a in enumerate(sys.argv):
        if a == '--mark':
            mark = sys.argv[i + 1]
        elif a == '--color':
            colors = {sys.argv[i + 1].upper()}
    sys.exit(main(args[0], args[1], mark, colors))
