#!/usr/bin/env python3
"""Chuyển docs/demo/demo-scenario.md sang xlsx nhiều sheet.

Không phải markdown-to-excel tổng quát: bám đúng cấu trúc của file kịch bản
(## màn hình -> > Note -> bảng test case) để xuất ra dạng bảng test case đọc
được trong Excel.
"""
from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path('/home/AI-Source-Code/agrinews/docs/demo/demo-scenario.md')
DST = Path('/home/AI-Source-Code/agrinews/docs/demo/demo-scenario.xlsx')

# ── màu ────────────────────────────────────────────────────────────────
C_SECTION = 'FF1F3864'   # nền heading màn hình
C_SUB = 'FFD9E2F3'       # nền heading mục con
C_NOTE = 'FFFFF2CC'      # nền khối Note
C_HEAD = 'FF4472C4'      # nền dòng tiêu đề cột
C_ALT = 'FFF2F2F2'       # nền xen kẽ
C_LINE = 'FFBFBFBF'

# Bản chữ của sơ đồ mermaid ở §0.4 — giữ nguyên nội dung, bỏ phần vẽ.
FLOW_TEXT = """\
(1) Bốn nguồn cùng GHI THÊM một dòng vào  t_dokusya_rireki  (PostgreSQL, append-only)

      SCR-011 購読者情報登録 — tạo / sửa / tái đặt   →  applyChange / insertResubscribe
      SCR-014 購読者明細検索 — nút 購読中止          →  insertScheduledKaiyaku
      batch dokusya-sync — 10 phút / lần            →  applyChange (source = BATCH)
            ↳ đọc cmsDB.users (collecting = 1) bên 電子版, khớp users.JACd → m_kanri_shiten
      batch dokusya-apply-due — 05:00 JST           →  insertKaiyaku (chốt hủy)

(2) t_dokusya_rireki  ──recomputeMaster──►  t_dokusya

      t_dokusya là ảnh chụp của dòng đang có hiệu lực HÔM NAY, không phải bảng bị sửa đè.

(3) Màn hình báo cáo đọc ở đâu

      t_dokusya — trạng thái   →  SCR-020 口座振替      SCR-021 配達手数料
      t_dokusya_rireki — biến động  →  SCR-026 購読者名簿    SCR-028 増減連絡票    SCR-029 増減通知"""

THIN = Side(style='thin', color=C_LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
WRAP_C = Alignment(wrap_text=True, vertical='center', horizontal='center')
BASE_FONT = 'Calibri'
MONO_FONT = 'Consolas'


# ── tiện ích ───────────────────────────────────────────────────────────
def inline(text: str) -> str:
    """Bỏ cú pháp markdown inline, giữ nguyên nội dung chữ."""
    text = re.sub(r'!\[([^\]]*)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text)
    text = text.replace('<br/>', '\n').replace('<br>', '\n')
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text, flags=re.S)
    text = re.sub(r'(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)', r'\1', text, flags=re.S)
    text = text.replace('`', '')
    text = re.sub(r'\\([*_`|])', r'\1', text)
    return text.strip()


def disp_width(s: str) -> int:
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in s)


def est_height(values, widths, min_h=18.0, pad=5.0, line_h=13.2, cap=409.0):
    lines = 1
    for val, w in zip(values, widths):
        if val is None:
            continue
        n = 0
        for seg in str(val).split('\n'):
            n += max(1, math.ceil(disp_width(seg) / max(4, w - 1.5)))
        lines = max(lines, n)
    return min(cap, max(min_h, lines * line_h + pad))


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip('|').split('|')]


def is_sep(line: str) -> bool:
    return bool(re.fullmatch(r'\|[\s:\-|]+\|', line.strip()))


# ── parse ──────────────────────────────────────────────────────────────
def parse(md: str):
    """-> list[(section, subsection, block)] với block = (kind, payload)."""
    lines = md.split('\n')
    out, i = [], 0
    section = subsection = None
    para: list[str] = []

    def flush_para():
        nonlocal para
        if para:
            txt = inline(' '.join(para).strip())
            if txt:
                out.append((section, subsection, ('para', txt)))
            para = []

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if s.startswith('# ') and not s.startswith('## '):
            flush_para()
            out.append((None, None, ('title', inline(s[2:]))))
            i += 1
        elif s.startswith('## '):
            flush_para()
            section, subsection = inline(s[3:]), None
            out.append((section, None, ('section', section)))
            i += 1
        elif s.startswith('### '):
            flush_para()
            subsection = inline(s[4:])
            out.append((section, subsection, ('sub', subsection)))
            i += 1
        elif s.startswith('---') and set(s) <= {'-'}:
            flush_para()
            i += 1
        elif s.startswith('>'):
            flush_para()
            buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(re.sub(r'^>\s?', '', lines[i].strip()))
                i += 1
            body = []
            for b in buf:
                if is_sep(b):
                    continue
                if b.startswith('|'):
                    body.append('   '.join(inline(c) for c in split_row(b)))
                else:
                    body.append(inline(b))
            out.append((section, subsection, ('note', '\n'.join(body).strip())))
        elif s.startswith('```'):
            flush_para()
            lang = s[3:].strip()
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            body = '\n'.join(buf).rstrip()
            if lang == 'mermaid':
                # Excel không vẽ được mermaid — thay bằng bản chữ cùng nội dung.
                body, lang = FLOW_TEXT, 'text'
            out.append((section, subsection, ('code', (lang, body))))
        elif s.startswith('|'):
            flush_para()
            head = split_row(s)
            i += 1
            if i < len(lines) and is_sep(lines[i]):
                i += 1
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([inline(c) for c in split_row(lines[i])])
                i += 1
            out.append((section, subsection, ('table', ([inline(h) for h in head], rows))))
        elif not s:
            flush_para()
            i += 1
        else:
            para.append(s)
            i += 1

    flush_para()
    return out


# ── ghi sheet ──────────────────────────────────────────────────────────
def setup(ws, widths):
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.sheet_view.showGridLines = False


def banner(ws, row, ncol, text, fill, color='FF000000', bold=True, size=11):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = PatternFill('solid', fgColor=fill)
    c.font = Font(name=BASE_FONT, bold=bold, size=size, color=color)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    return c


def write_free(ws, blocks, widths, ncol):
    """Sheet dạng tài liệu: đoạn văn / bảng / code xếp nối tiếp."""
    setup(ws, widths)
    r = 1
    for _, _, (kind, payload) in blocks:
        if kind in ('title', 'section'):
            banner(ws, r, ncol, payload, C_SECTION, color='FFFFFFFF', size=13)
            ws.row_dimensions[r].height = 24
            r += 2
        elif kind == 'sub':
            banner(ws, r, ncol, payload, C_SUB, size=11)
            ws.row_dimensions[r].height = 20
            r += 1
        elif kind == 'para':
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
            c = ws.cell(row=r, column=1, value=payload)
            c.alignment = WRAP
            c.font = Font(name=BASE_FONT, size=10)
            ws.row_dimensions[r].height = est_height([payload], [sum(widths)])
            r += 2
        elif kind == 'note':
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
            c = ws.cell(row=r, column=1, value=payload)
            c.fill = PatternFill('solid', fgColor=C_NOTE)
            c.alignment = WRAP
            c.font = Font(name=BASE_FONT, size=10)
            c.border = BORDER
            ws.row_dimensions[r].height = est_height([payload], [sum(widths)])
            r += 2
        elif kind == 'code':
            _, body = payload
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
            c = ws.cell(row=r, column=1, value=body)
            c.alignment = WRAP
            c.font = Font(name=MONO_FONT, size=9)
            c.fill = PatternFill('solid', fgColor='FFF7F7F7')
            c.border = BORDER
            ws.row_dimensions[r].height = est_height([body], [sum(widths)], line_h=12.0)
            r += 2
        elif kind == 'table':
            head, rows = payload
            n = len(head)
            for j, h in enumerate(head, start=1):
                c = ws.cell(row=r, column=j, value=h)
                c.fill = PatternFill('solid', fgColor=C_HEAD)
                c.font = Font(name=BASE_FONT, bold=True, size=10, color='FFFFFFFF')
                c.alignment = WRAP_C
                c.border = BORDER
            ws.row_dimensions[r].height = 22
            r += 1
            for k, row in enumerate(rows):
                for j in range(n):
                    val = row[j] if j < len(row) else ''
                    c = ws.cell(row=r, column=j + 1, value=val)
                    c.alignment = WRAP
                    c.font = Font(name=BASE_FONT, size=10)
                    c.border = BORDER
                    if k % 2:
                        c.fill = PatternFill('solid', fgColor=C_ALT)
                ws.row_dimensions[r].height = est_height(row, widths[:n])
                r += 1
            r += 1
    return r


TC_HEAD = ['Scenario ID', 'Scenario Name', 'Pre-condition', 'Steps', 'Expected Result']
TC_WIDTHS = [13, 42, 34, 58, 72]


def write_cases(ws, blocks):
    setup(ws, TC_WIDTHS)
    ncol = len(TC_HEAD)
    for j, h in enumerate(TC_HEAD, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = PatternFill('solid', fgColor=C_HEAD)
        c.font = Font(name=BASE_FONT, bold=True, size=10, color='FFFFFFFF')
        c.alignment = WRAP_C
        c.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    r, n_tc = 2, 0
    for _, _, (kind, payload) in blocks:
        if kind == 'section':
            banner(ws, r, ncol, payload, C_SECTION, color='FFFFFFFF', size=12)
            ws.row_dimensions[r].height = 24
            r += 1
        elif kind == 'sub':
            banner(ws, r, ncol, payload, C_SUB, size=11)
            ws.row_dimensions[r].height = 20
            r += 1
        elif kind in ('note', 'para'):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
            c = ws.cell(row=r, column=1, value=payload)
            c.fill = PatternFill('solid', fgColor=C_NOTE)
            c.alignment = WRAP
            c.font = Font(name=BASE_FONT, size=10)
            c.border = BORDER
            ws.row_dimensions[r].height = est_height([payload], [sum(TC_WIDTHS)])
            r += 1
        elif kind == 'table':
            head, rows = payload
            if head[:1] != ['Scenario ID']:
                continue
            for row in rows:
                for j in range(ncol):
                    val = row[j] if j < len(row) else ''
                    c = ws.cell(row=r, column=j + 1, value=val)
                    c.alignment = WRAP
                    c.font = Font(name=BASE_FONT, size=10,
                                  bold=(j == 0))
                    c.border = BORDER
                ws.cell(row=r, column=1).alignment = WRAP_C
                ws.row_dimensions[r].height = est_height(row, TC_WIDTHS)
                r += 1
                n_tc += 1
    return n_tc


def main() -> None:
    blocks = parse(SRC.read_text(encoding='utf-8'))

    def sec_of(b):
        return b[0]

    overview = [b for b in blocks if sec_of(b) is None]
    prep = [b for b in blocks if sec_of(b) and sec_of(b).startswith('0.')]
    cases = [b for b in blocks if sec_of(b)
             and (sec_of(b).startswith('Màn hình') or sec_of(b).startswith('Batch'))]
    apx = {}
    for b in blocks:
        s = sec_of(b) or ''
        if s.startswith('Phụ lục'):
            apx.setdefault(s, []).append(b)

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('概要 Tổng quan')
    ov_w = [24, 24, 24, 24, 24, 24]
    r = write_free(ws, overview, ov_w, 6)
    toc = [
        ['Sheet', 'Nội dung'],
        ['事前準備 Chuẩn bị',
         'Lệnh dựng dữ liệu, bảng ngày phải gõ vào từng màn hình, master data có sẵn, '
         'toàn cảnh luồng dữ liệu. Làm 1 lần trước khi demo, không tính là test case.'],
        ['テストケース',
         'TC-001 → TC-031 chạy tuần tự, nhóm theo màn hình: SCR-011 tạo/sửa → batch '
         'dokusya-sync → SCR-014 hủy 2 pha → SCR-011 tái đặt → 5 màn báo cáo '
         '(SCR-020/021/026/028/029). Dòng nền vàng là điều kiện/ghi chú của cả nhóm.'],
        ['付録A 登場人物', '12 độc giả demo và vai trò của từng người, kèm việc họ có lên từng báo cáo hay không.'],
        ['付録B トラブル対応', 'Hiện tượng → nguyên nhân → cách xử lý khi số liệu không khớp.'],
        ['付録C クリーンアップ', 'Lệnh xoá dữ liệu demo ở cả cloud và mock cmsDB.'],
    ]
    banner(ws, r, 6, 'Hướng dẫn đọc file', C_SUB, size=11)
    ws.row_dimensions[r].height = 20
    r += 1
    for k, row in enumerate(toc):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=val)
            c.alignment = WRAP
            c.border = BORDER
            c.font = Font(name=BASE_FONT, size=10, bold=(k == 0),
                          color='FFFFFFFF' if k == 0 else 'FF000000')
            if k == 0:
                c.fill = PatternFill('solid', fgColor=C_HEAD)
            elif k % 2 == 0:
                c.fill = PatternFill('solid', fgColor=C_ALT)
        for j in range(2, 7):
            ws.cell(row=r, column=j).border = BORDER
        ws.row_dimensions[r].height = est_height(row, [ov_w[0], sum(ov_w[1:])])
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1,
                value='File này sinh từ docs/demo/demo-scenario.md — sửa bản .md rồi tạo lại, '
                      'đừng sửa trực tiếp .xlsx.')
    c.font = Font(name=BASE_FONT, size=9, italic=True, color='FF808080')
    c.alignment = WRAP

    ws = wb.create_sheet('事前準備 Chuẩn bị')
    write_free(ws, prep, [30, 30, 30, 30], 4)

    ws = wb.create_sheet('テストケース')
    n_tc = write_cases(ws, cases)

    names = {'A': '付録A 登場人物', 'B': '付録B トラブル対応', 'C': '付録C クリーンアップ'}
    widths = {'A': [9, 14, 10, 46, 8, 8, 8, 20],
              'B': [40, 46, 52],
              'C': [30, 30, 30]}
    for key, blks in apx.items():
        letter = key.split()[2] if len(key.split()) > 2 else key[-1]
        letter = letter.strip('—').strip() or key[-1]
        letter = letter[0]
        ws = wb.create_sheet(names.get(letter, key[:31]))
        w = widths.get(letter, [30] * 4)
        write_free(ws, blks, w, len(w))

    for s in wb.worksheets:
        s.page_setup.orientation = 'landscape'
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(DST)
    print(f'sheets  : {[s.title for s in wb.worksheets]}')
    print(f'test-case rows : {n_tc}')
    print(f'saved   : {DST}  ({DST.stat().st_size:,} bytes)')


if __name__ == '__main__':
    main()
